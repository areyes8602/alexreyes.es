#!/usr/bin/env python3
"""Carga los grupos de nivel de matemáticas de 2º de ESO en D1.

Lee el Excel que reparte los cinco grupos-clase de 2º en seis grupos de
nivel (Alt, Mig 1, Mig 2, Mig 3, Baix 1, Baix 2), cada uno con su profesor.

IMPORTANTE — dónde acaban los datos
-----------------------------------
Son datos de menores. Como el resto de importadores de tutoría, este NO
escribe dentro del repositorio, que es público: el destino es D1 y el SQL
intermedio va a un directorio temporal que debes borrar al subirlo.

Uso
---
    python3 scripts/mates_import_nivells.py --excel Nivells_2n.xlsx [--dry-run]

Antes, una sola vez:
    npx wrangler d1 execute tutoria --remote --file=scripts/sql/mates_schema.sql

Cómo cruza con tutoría
----------------------
La clave es el mismo slug que usa tutoria_import_orla.py, slug("Cognoms
Nom"). El Excel trae "Cognoms, Nom" en una sola casilla y se parte por la
coma. Si algún nombre no cuadrara con el de la orla, la ficha de tutoría no
encontraría su grupo de mates: por eso el script cuenta cuántos ha leído de
cada clase, para que veas si falta alguien.

El aula NO viene en el Excel y no se toca: la vas poniendo tú desde
/mates/, y volver a pasar el importador no la borra.
"""
import argparse
import re
import sys
import tempfile
import unicodedata
from collections import Counter
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent

# La hoja "Hoja1" es la tabla limpia: una fila por alumno y sin columnas
# repetidas. Las otras dos ("Grups_2n", "Llistats") son la misma gente
# presentada de otra manera.
FULL = "Hoja1"
COLS = ("Grup", "Cognoms i nom", "Ve de", "Professor")

# De más a menos, que es como se miran, no por orden alfabético.
ORDRE = ["Alt", "Mig 1", "Mig 2", "Mig 3", "Baix 1", "Baix 2"]


def slug(text):
    t = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    t = re.sub(r"[^a-zA-Z0-9]+", "-", t).strip("-").lower()
    return re.sub(r"-{2,}", "-", t)


def sql_str(v):
    if v is None or str(v).strip() == "":
        return "NULL"
    return "'" + str(v).strip().replace("'", "''") + "'"


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--excel", required=True)
    ap.add_argument("--curs", default="2026-27")
    ap.add_argument("--out")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("Hace falta openpyxl:  pip3 install openpyxl")

    wb = load_workbook(args.excel, data_only=True)
    if FULL not in wb.sheetnames:
        sys.exit(f"No encuentro la hoja «{FULL}». El fichero tiene: {wb.sheetnames}")
    ws = wb[FULL]

    caps = [str(c or "").strip() for c in next(ws.iter_rows(max_row=1, values_only=True))]
    if caps[:4] != list(COLS):
        sys.exit(f"Las columnas no son las esperadas.\n  esperaba: {list(COLS)}\n  encontré: {caps[:4]}")

    alumnes, dolents = [], []
    for n, fila in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        nivell, qui, origen, prof = [str(c).strip() if c is not None else "" for c in fila[:4]]
        if not nivell or not qui:
            continue
        if "," in qui:
            cognoms, nom = [p.strip() for p in qui.split(",", 1)]
        else:                      # cognom sol, sense nom de pila
            cognoms, nom = qui, ""
        if not cognoms:
            dolents.append(n)
            continue
        alumnes.append({"id": slug(f"{cognoms} {nom}".strip()), "nivell": nivell,
                        "origen": origen, "cognoms": cognoms, "nom": nom, "prof": prof})

    print(f"Alumnos leídos : {len(alumnes)}")
    if dolents:
        print(f"  filas que no he sabido leer: {dolents}")
    print("  por nivel     :", dict(Counter(a["nivell"] for a in alumnes)))
    print("  por clase     :", dict(Counter(a["origen"] for a in alumnes)))

    # Un nivel con dos profesores distintos es un error del Excel, no algo
    # que deba acabar en la base callando.
    grups = {}
    for a in alumnes:
        grups.setdefault(a["nivell"], set()).add(a["prof"])
    mal = {n: p for n, p in grups.items() if len(p) > 1}
    if mal:
        sys.exit(f"Un mismo nivel con varios profesores: { {n: sorted(p) for n, p in mal.items()} }")

    repes = [i for i, c in Counter(a["id"] for a in alumnes).items() if c > 1]
    if repes:
        print(f"  ¡ojo! {len(repes)} alumnos aparecen más de una vez; se cargará el último.")

    if args.dry_run:
        print("\n--dry-run: no se ha escrito nada.")
        return

    out = Path(args.out) if args.out else Path(tempfile.mkdtemp(prefix="mates-"))
    if REPO in out.resolve().parents or out.resolve() == REPO:
        sys.exit(f"El directorio de salida está dentro del repositorio ({out}).\n"
                 "Elige uno fuera: estos datos no deben acabar en git.")
    out.mkdir(parents=True, exist_ok=True)

    c = sql_str(args.curs)
    linies = ["-- Generado por scripts/mates_import_nivells.py — NO commitear.",
              "-- Datos de menores. Bórralo en cuanto lo hayas subido."]
    for nivell in sorted(grups, key=lambda n: ORDRE.index(n) if n in ORDRE else 99):
        prof = sorted(grups[nivell])[0]
        # El aula no se toca: la pone el profesor desde la página.
        linies.append(
            f"INSERT INTO mates_grups (curs, nivell, professor, ordre, updated_at) "
            f"VALUES ({c}, {sql_str(nivell)}, {sql_str(prof)}, "
            f"{ORDRE.index(nivell) if nivell in ORDRE else 99}, datetime('now')) "
            f"ON CONFLICT(curs, nivell) DO UPDATE SET professor = excluded.professor, "
            f"ordre = excluded.ordre, updated_at = excluded.updated_at;")
    for a in alumnes:
        linies.append(
            f"INSERT INTO mates_alumnes (curs, id, nivell, grup_origen, cognoms, nom, updated_at) "
            f"VALUES ({c}, {sql_str(a['id'])}, {sql_str(a['nivell'])}, {sql_str(a['origen'])}, "
            f"{sql_str(a['cognoms'])}, {sql_str(a['nom'])}, datetime('now')) "
            f"ON CONFLICT(curs, id) DO UPDATE SET nivell = excluded.nivell, "
            f"grup_origen = excluded.grup_origen, cognoms = excluded.cognoms, "
            f"nom = excluded.nom, updated_at = excluded.updated_at;")
    fitxer = out / "mates-nivells.sql"
    fitxer.write_text("\n".join(linies) + "\n", encoding="utf-8")

    print(f"\nSQL escrito en:\n  {fitxer}\n")
    print("Súbelo y borra el rastro:")
    print(f"  npx wrangler d1 execute tutoria --remote --file={fitxer}")
    print(f"  rm -rf {out}")


if __name__ == "__main__":
    main()
